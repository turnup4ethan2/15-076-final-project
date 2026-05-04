"""Build a Word doc with speaker scripts for slides 6, 13, 14, 15."""

from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches

OUT = Path(__file__).resolve().parent / "presentation_scripts.docx"

SLIDES = [
    (
        "Slide 6 — Quantile LSTM Model: Forecasting Reservoir Supply",
        [
            "This is the first of three stages — forecasting how much each reservoir's storage will change one to three months ahead, individually for every reservoir in the basin.",
            "We chose an LSTM — Long Short-Term Memory network — because reservoir levels are strongly seasonal and depend on history that runs months back. They fill in winter, drain in summer, and the relevant context isn't just last month — it's the whole prior year. LSTMs are built exactly for that kind of sequential structure; a feedforward model would lose the temporal context.",
            "The key design choice is on the output side. Instead of predicting a single point estimate, we output three quantiles per reservoir per month — a pessimistic P10, a middle P50, and an optimistic P90. We train this with pinball loss, the standard objective for quantile regression. That uncertainty range is what lets the downstream optimizer plan robustly, rather than betting everything on a single number.",
            "The table on the right is our calibration check. We trained on data from 1988 through 2022 and held out 2023 onward. When we say 'P10,' the true outcome should fall below it about 10 percent of the time. And across all three forecast horizons — one, two, and three months out — our quantiles land within a couple of points of target. The model is well-calibrated, which means the optimizer can take these forecasts at face value.",
        ],
    ),
    (
        "Slide 13 — Results: Water Allocation (P10 baseline)",
        [
            "This is the optimizer's output under our worst-case supply forecast.",
            "Four of the five provinces — Jaén, Córdoba, Granada, and Huelva — get their full demand met. The interesting case is Sevilla. Sevilla has the largest demand of any province in the basin, around 1,466 hm³ a year, but the LP only delivers 484 — about 33 percent.",
            "That's not a forecasting failure. It's structural. The reservoirs we route to Sevilla physically can't store or supply enough water to cover the rice belt under drought conditions. In real operations, the basin authority bridges this gap with inter-district transfers and groundwater pumping — neither of which we model. So this shortfall is the LP surfacing a real geographic limit of the basin.",
        ],
    ),
    (
        "Slide 14 — Results: Changing the supply assumption",
        [
            "To check how much our results depend on the forecast, we re-ran the optimization three times with progressively less pessimistic supply — strict, moderate, and mild.",
            "The finding is striking. Four of the five provinces get the same allocation in every scenario. The only thing that moves is Sevilla, climbing from 484 to 524 to 563 hm³ as the supply assumption eases.",
            "The reason is what we just saw: the other four provinces are already pinned at their structural ceilings, either their demand cap or the reservoir-routing limit. Sevilla is the only province with unmet demand, so every additional hm³ of preserved storage flows to it. The LP isn't really choosing between provinces here — the geography has already decided four of five, and Sevilla is the only lever left.",
        ],
    ),
    (
        "Slide 15 — Limitations & Out of Scope",
        [
            "Six things we want to be honest about.",
            "First, granularity — we allocate to whole provinces, not the irrigation districts water actually flows through. Second, ecological flows — real reservoirs have minimum-release requirements that we don't enforce. Third, our demand model uses static crop-water coefficients; it's the same every year because we deferred AEMET weather integration.",
            "Fourth, our LSTM forecasts net change in storage rather than gauged inflow. Fifth, we treat reservoirs as standalone — no groundwater, no hydropower coupling — even though Spanish water management leans heavily on both. And sixth, the LP doesn't enforce institutional water rights, so its outputs are exploratory recommendations, not deployment-ready policy.",
            "The clearest next steps are AEMET integration, ecological flow constraints, and finer-grain district modeling.",
        ],
    ),
]


def main() -> None:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    title = doc.add_heading("15.076 Final Presentation — Speaker Scripts", level=0)
    sub = doc.add_paragraph()
    sub.add_run("Guadalquivir Water Allocation Optimization · Spring 2026").italic = True

    for heading, paragraphs in SLIDES:
        doc.add_heading(heading, level=1)
        for p_text in paragraphs:
            p = doc.add_paragraph(p_text)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.4

    doc.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
